"""
Parser fuer die bisherige Excel-Buchfuehrung (Issue #16).

Aufbau der Arbeitsmappe (Kopfzeilen werden am Text erkannt, nicht an festen
Zelladressen):

* "Kraftstoffkosten LPG" / "Kraftstoffkosten Benzin": oben eine
  Monatsuebersicht, darunter die eigentliche Tabelle mit einer Zeile je
  Tankung (Datum, Kilometerstand, Füllmenge, Preis/ Liter, Gesamtpreis,
  ..., nicht voll). Die Kraftstoffart ergibt sich aus dem Blattnamen. Die
  erste Zeile ohne Füllmenge ist der Start-Kilometerstand beim Kauf, ein
  "X" in der Datumsspalte markiert das Tabellenende.
* "Gesamtstatistik": nur abgeleitete Daten (Matrixformeln, die die
  Tankblaetter und sonstigen Kosten zusammenfuehren, ohne Füllmenge und
  ohne Kraftstoffart). Wird bewusst NICHT importiert - sonst kaemen alle
  Tankungen doppelt.
* "Sonstige Kosten": Kopfzeile "Datum", dann je Kategorie eine Betragsspalte
  (Einnahmen RKA, Versicherung, Steuer, ...), danach berechnete Spalten
  ("... bis heute", "Unterhaltskosten", "Gesamtkosten") und eine
  Beschreibungsspalte ohne Ueberschrift. Die Zeile "Erwerb" ist der
  Fahrzeugkauf und wird als Kaufdaten-Vorschlag gemeldet statt als Kosten.

Die Funktionen hier sind rein (keine DB), damit sie einzeln testbar sind.
"""
from __future__ import annotations

import datetime
import io
import re
import zipfile
from dataclasses import dataclass, field

from app.models import Kostenkategorie, Kraftstoffart

MAX_ROWS = 10000
MAX_COLS = 60
MAX_WARNUNGEN = 200
MAX_UNCOMPRESSED_BYTES = 400 * 1024 * 1024


class ExcelImportError(ValueError):
    """Datei ist keine lesbare Arbeitsmappe oder enthaelt nichts Importierbares."""


@dataclass
class ImportFuelEntry:
    datum: datetime.date
    kilometerstand: int
    kraftstoffart: Kraftstoffart
    fuellmenge_liter: float
    preis_pro_liter: float
    gesamtpreis: float
    nicht_voll: bool
    quelle: str  # z.B. "Kraftstoffkosten LPG, Zeile 20"


@dataclass
class ImportOtherCost:
    datum: datetime.date
    kategorie: Kostenkategorie
    betrag: float
    beschreibung: str | None
    jaehrlich_wiederkehrend: bool
    quelle: str


@dataclass
class KaufVorschlag:
    kaufdatum: datetime.date | None = None
    kaufpreis: float | None = None
    kaufkilometerstand: int | None = None
    quelle: str | None = None


@dataclass
class ImportPreview:
    fuel_entries: list[ImportFuelEntry] = field(default_factory=list)
    other_costs: list[ImportOtherCost] = field(default_factory=list)
    kauf: KaufVorschlag | None = None
    warnungen: list[str] = field(default_factory=list)

    def warn(self, message: str) -> None:
        if len(self.warnungen) < MAX_WARNUNGEN:
            self.warnungen.append(message)
        elif len(self.warnungen) == MAX_WARNUNGEN:
            self.warnungen.append("Weitere Hinweise ausgeblendet.")


# --- Hilfsfunktionen fuer Zellwerte ----------------------------------------


def _norm(value: object) -> str:
    """Kopfzeilentext vergleichbar machen: 'Preis/ Liter' -> 'preisliter'."""
    text = str(value or "").lower()
    for src, dst in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        text = text.replace(src, dst)
    return re.sub(r"[^a-z0-9]", "", text)


def _to_float(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace("€", "").replace(" ", "")
        if not text or text.startswith("#"):  # leer oder Excel-Fehler wie #N/A
            return None
        if "," in text:  # deutsches Zahlenformat 1.234,56
            text = text.replace(".", "").replace(",", ".")
        try:
            return float(text)
        except ValueError:
            return None
    return None


def _to_date(value: object) -> datetime.date | None:
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 20000 < value < 80000:
        # Unformatierte Excel-Seriennummer (Tage seit 1899-12-30)
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(value))
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
            try:
                return datetime.datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def _truthy(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return _norm(value) not in ("", "0", "nein", "no", "false", "falsch")


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _datum_fehlt(raw: object) -> str:
    return "kein Datum" if _is_blank(raw) else f"Datum „{raw}“ nicht lesbar"


def _cell(row: tuple, idx: int | None) -> object:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


# --- Tankblaetter ------------------------------------------------------------

_FUEL_COLUMNS = {
    "datum": lambda h: h == "datum" or h.startswith("tankdatum"),
    "km": lambda h: h.startswith("kilometerstand") or h in ("kmstand", "km", "tacho"),
    "liter": lambda h: h in ("fuellmenge", "fuellmengel", "menge", "mengel", "liter"),
    "preis": lambda h: h in ("preisliter", "preisproliter", "literpreis", "preisl"),
    "gesamt": lambda h: h in ("gesamtpreis", "betrag", "summe", "preisgesamt"),
    "nicht_voll": lambda h: h in ("nichtvoll", "teilbetankung", "teilgetankt"),
    "kraftstoff": lambda h: h in ("kraftstoff", "kraftstoffart", "sorte"),
}


def _find_fuel_header(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    for r_idx, row in enumerate(rows):
        cols: dict[str, int] = {}
        for c_idx, value in enumerate(row):
            if not isinstance(value, str):
                continue
            h = _norm(value)
            for key, match in _FUEL_COLUMNS.items():
                if key not in cols and match(h):
                    cols[key] = c_idx
                    break
        if {"datum", "km", "liter"} <= cols.keys() and ("preis" in cols or "gesamt" in cols):
            return r_idx, cols
    return None


def _fuel_type(text: object) -> Kraftstoffart | None:
    h = _norm(text)
    if "lpg" in h or "autogas" in h:
        return Kraftstoffart.LPG
    if "benzin" in h or "super" in h or "e10" in h:
        return Kraftstoffart.BENZIN
    if "diesel" in h:
        return Kraftstoffart.DIESEL
    if "strom" in h or "kwh" in h:
        return Kraftstoffart.STROM
    return None


def _parse_fuel_sheet(title: str, rows: list[tuple], header: tuple[int, dict[str, int]], preview: ImportPreview,
                      start_punkte: list[tuple[datetime.date, int, str]]) -> None:
    header_idx, cols = header
    sheet_type = _fuel_type(title)
    if sheet_type is None and "kraftstoff" not in cols:
        preview.warn(f"Blatt „{title}“: Kraftstoffart nicht erkennbar (Blattname ohne LPG/Benzin/Diesel) – übersprungen.")
        return

    tankungen_bisher = 0
    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        where = f"{title}, Zeile {offset}"
        km = _to_float(_cell(row, cols["km"]))
        liter = _to_float(_cell(row, cols["liter"]))
        preis = _to_float(_cell(row, cols.get("preis")))
        gesamt = _to_float(_cell(row, cols.get("gesamt")))
        raw_datum = _cell(row, cols["datum"])
        if all(v is None for v in (km, liter, preis, gesamt)):
            continue  # Leerzeile oder Endmarkierung "X"

        datum = _to_date(raw_datum)
        if datum is None:
            preview.warn(f"{where}: {_datum_fehlt(raw_datum)} – übersprungen.")
            continue

        if not liter and not gesamt:
            if tankungen_bisher == 0 and km:
                # Startzeile (Kauf-km-Stand), ab der der erste Verbrauch rechnet
                start_punkte.append((datum, int(round(km)), where))
            else:
                preview.warn(f"{where}: ohne Füllmenge/Betrag – übersprungen.")
            continue

        # fehlenden dritten Wert aus den beiden anderen ergaenzen
        if gesamt is None and liter is not None and preis is not None:
            gesamt = liter * preis
        elif liter is None and gesamt is not None and preis:
            liter = gesamt / preis
        elif preis is None and gesamt is not None and liter:
            preis = gesamt / liter

        if km is None:
            preview.warn(f"{where}: kein Kilometerstand – übersprungen.")
            continue
        if not liter or not gesamt or not preis or liter <= 0 or gesamt <= 0 or preis <= 0:
            preview.warn(f"{where}: Füllmenge, Preis/Liter oder Gesamtpreis fehlt oder ist 0 – übersprungen.")
            continue

        art = _fuel_type(_cell(row, cols.get("kraftstoff"))) or sheet_type
        if art is None:
            preview.warn(f"{where}: Kraftstoffart nicht erkennbar – übersprungen.")
            continue

        tankungen_bisher += 1
        preview.fuel_entries.append(
            ImportFuelEntry(
                datum=datum,
                kilometerstand=int(round(km)),
                kraftstoffart=art,
                fuellmenge_liter=round(liter, 3),
                preis_pro_liter=round(preis, 4),
                gesamtpreis=round(gesamt, 2),
                nicht_voll=_truthy(_cell(row, cols.get("nicht_voll"))),
                quelle=where,
            )
        )


# --- Sonstige Kosten ---------------------------------------------------------

_KAUF_PATTERN = re.compile(r"\b(erwerb|kauf|kaufpreis|fahrzeugkauf|autokauf)\b", re.IGNORECASE)


def _cost_column_role(h: str) -> tuple[str, Kostenkategorie | None]:
    """Rolle einer Kopfzeilenspalte im Blatt 'Sonstige Kosten'."""
    if not h:
        return "leer", None
    if "bisheute" in h or h.startswith("gesamt") or h.startswith("unterhalt") or "protag" in h or h.startswith("summe"):
        return "berechnet", None
    if h.startswith("jaehrlich"):
        return "jaehrlich", None
    if h.startswith(("beschreibung", "bemerkung", "notiz", "kommentar", "text", "verwendungszweck")):
        return "beschreibung", None
    if h.startswith("einnahme"):
        return "kategorie", Kostenkategorie.EINNAHME
    if h.startswith("versicherung"):
        return "kategorie", Kostenkategorie.VERSICHERUNG
    if h.startswith("steuer"):
        return "kategorie", Kostenkategorie.STEUER
    if h.startswith(("finanzierung", "kredit", "leasing", "rate")):
        return "kategorie", Kostenkategorie.FINANZIERUNG
    if "reifen" in h or "teile" in h or "zubehoer" in h:
        return "kategorie", Kostenkategorie.REIFEN_TEILE
    if "service" in h or "tuev" in h or "inspektion" in h or "werkstatt" in h or "reparatur" in h:
        return "kategorie", Kostenkategorie.SERVICE_TUEV
    if "wasch" in h or "waesche" in h or "pflege" in h:
        return "kategorie", Kostenkategorie.WASCHEN
    if h.startswith(("einmalig", "anschaffung")):
        return "kategorie", Kostenkategorie.EINMALIG
    if h.startswith("sonstig"):
        return "kategorie", Kostenkategorie.SONSTIGES
    return "unbekannt", Kostenkategorie.SONSTIGES


def _find_cost_header(rows: list[tuple]) -> tuple[int, int, dict] | None:
    """Liefert (Zeilenindex, Datumsspalte, Spaltenrollen) der Kopfzeile."""
    for r_idx, row in enumerate(rows):
        date_col = next(
            (c for c, v in enumerate(row) if isinstance(v, str) and _norm(v).startswith("datum")), None
        )
        if date_col is None:
            continue
        roles: dict[int, tuple[str, Kostenkategorie | None, str]] = {}
        block_offen = True  # unbekannte Spalten zaehlen nur vor den berechneten
        for c_idx in range(date_col + 1, len(row)):
            label = str(row[c_idx]).strip() if not _is_blank(row[c_idx]) else ""
            role, kategorie = _cost_column_role(_norm(label))
            if role == "berechnet":
                block_offen = False
            if role == "unbekannt" and not block_offen:
                role, kategorie = "ignoriert", None
            roles[c_idx] = (role, kategorie, label)
        if any(role == "kategorie" for role, _, _ in roles.values()):
            return r_idx, date_col, roles
    return None


def _parse_cost_sheet(title: str, rows: list[tuple], header: tuple[int, int, dict], preview: ImportPreview) -> None:
    header_idx, date_col, roles = header
    category_cols = {c: (k, label) for c, (role, k, label) in roles.items() if role in ("kategorie", "unbekannt")}
    flag_col = next((c for c, (role, _, _) in roles.items() if role == "jaehrlich"), None)
    desc_cols = [c for c, (role, _, _) in roles.items() if role == "beschreibung"]
    # Ohne eigene Ueberschrift: erste Textzelle in einer leeren/ignorierten Spalte
    fallback_desc_cols = [c for c, (role, _, _) in roles.items() if role in ("leer", "ignoriert")]

    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        where = f"{title}, Zeile {offset}"
        amounts: list[tuple[Kostenkategorie, str, float]] = []
        for c_idx, (kategorie, label) in category_cols.items():
            value = _cell(row, c_idx)
            if _is_blank(value):
                continue
            betrag = _to_float(value)
            if betrag is None:
                preview.warn(f"{where}: Betrag „{value}“ in Spalte „{label}“ nicht lesbar – übersprungen.")
            elif abs(betrag) >= 0.005:
                amounts.append((kategorie, label, betrag))
        if not amounts:
            continue

        raw_datum = _cell(row, date_col)
        datum = _to_date(raw_datum)
        if datum is None:
            preview.warn(f"{where}: {_datum_fehlt(raw_datum)} – übersprungen.")
            continue

        beschreibung = next(
            (str(_cell(row, c)).strip() for c in desc_cols + fallback_desc_cols
             if isinstance(_cell(row, c), str) and _cell(row, c).strip()),
            None,
        )
        jaehrlich = _truthy(_cell(row, flag_col)) if flag_col is not None else False

        if preview.kauf is None and beschreibung and _KAUF_PATTERN.search(beschreibung):
            preview.kauf = KaufVorschlag(
                kaufdatum=datum, kaufpreis=round(sum(b for _, _, b in amounts), 2), quelle=where
            )
            continue

        for kategorie, label, betrag in amounts:
            text = beschreibung
            if _label_geht_verloren(kategorie, label):
                text = f"{label}: {beschreibung}" if beschreibung else label
            preview.other_costs.append(
                ImportOtherCost(
                    datum=datum,
                    kategorie=kategorie,
                    betrag=round(betrag, 2),
                    beschreibung=text[:300] if text else None,
                    jaehrlich_wiederkehrend=jaehrlich,
                    quelle=where,
                )
            )


def _label_geht_verloren(kategorie: Kostenkategorie, label: str) -> bool:
    """Originale Spaltenbezeichnung behalten, wenn die Kategorie sie nicht
    wiedergibt (Auffangkategorie oder z.B. 'Einnahmen RKA' -> Einnahme)."""
    return kategorie in (Kostenkategorie.SONSTIGES, Kostenkategorie.EINNAHME) and _norm(label) not in (
        _norm(kategorie.value),
        _norm(kategorie.value) + "n",
    )


# --- Einstieg ----------------------------------------------------------------


def _read_rows(ws) -> list[tuple]:
    rows = []
    for row in ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True):
        rows.append(tuple(row))
    # abschliessende Leerzeilen abschneiden
    while rows and all(_is_blank(v) for v in rows[-1]):
        rows.pop()
    return rows


def _check_zip_size(file_bytes: bytes) -> None:
    """Schutz vor "Zip-Bomben": .xlsx ist ein ZIP-Archiv, dessen entpackte
    Groesse vor dem Parsen begrenzt wird (die echte Arbeitsmappe hat ~65 MB
    entpackt, vor allem wegen vieler Formelzellen)."""
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ExcelImportError(
            "Die Datei konnte nicht als Excel-Arbeitsmappe (.xlsx) gelesen werden."
        ) from exc
    if total > MAX_UNCOMPRESSED_BYTES:
        raise ExcelImportError("Die Arbeitsmappe ist entpackt zu groß für den Import.")


def parse_workbook(file_bytes: bytes) -> ImportPreview:
    _check_zip_size(file_bytes)
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - zip/xml/Formatfehler aller Art
        raise ExcelImportError(
            "Die Datei konnte nicht als Excel-Arbeitsmappe (.xlsx) gelesen werden."
        ) from exc

    preview = ImportPreview()
    start_punkte: list[tuple[datetime.date, int, str]] = []
    fuel_sheets = cost_sheets = 0
    try:
        for ws in wb.worksheets:
            title = ws.title
            if "gesamtstatistik" in _norm(title):
                continue  # nur abgeleitete Daten, siehe Moduldoku
            rows = _read_rows(ws)
            fuel_header = _find_fuel_header(rows)
            if fuel_header is not None:
                fuel_sheets += 1
                _parse_fuel_sheet(title, rows, fuel_header, preview, start_punkte)
                continue
            cost_header = _find_cost_header(rows)
            if cost_header is not None:
                cost_sheets += 1
                _parse_cost_sheet(title, rows, cost_header, preview)
    finally:
        wb.close()

    if not fuel_sheets and not cost_sheets:
        raise ExcelImportError(
            "Keine Tank- oder Kostentabelle gefunden (erwartet z.B. Blätter "
            "„Kraftstoffkosten LPG“ und „Sonstige Kosten“ mit Kopfzeile „Datum“)."
        )

    if start_punkte:
        datum, km, where = min(start_punkte, key=lambda p: p[1])
        if preview.kauf is None:
            preview.kauf = KaufVorschlag(quelle=where)
        preview.kauf.kaufkilometerstand = km
        if preview.kauf.kaufdatum is None:
            preview.kauf.kaufdatum = datum

    preview.fuel_entries.sort(key=lambda e: (e.datum, e.kilometerstand))
    preview.other_costs.sort(key=lambda c: c.datum)
    return preview


def km_ausreisser(punkte: list[tuple[datetime.date, int, bool]]) -> set[int]:
    """Indizes neuer Punkte, die nicht in die Kilometerfolge passen.

    punkte: (datum, km, ist_neu). Gesucht wird die laengste chronologische
    Folge mit nicht sinkendem km-Stand; vorhandene Eintraege wiegen dabei so
    schwer, dass sie praktisch immer drinbleiben. Was von den neuen Punkten
    nicht in diese Folge passt, ist vermutlich ein Tippfehler (z.B. 181.901
    statt 172.901) und wird beim Import uebersprungen statt die ganze Datei
    abzulehnen.
    """
    reihenfolge = sorted(range(len(punkte)), key=lambda i: (punkte[i][0], punkte[i][1]))
    gewicht = [1 if neu else 10_000 for _, _, neu in punkte]
    beste: dict[int, int] = {}
    vorgaenger: dict[int, int | None] = {}
    for pos, i in enumerate(reihenfolge):
        beste[i], vorgaenger[i] = gewicht[i], None
        for j in reihenfolge[:pos]:
            if punkte[j][1] <= punkte[i][1] and beste[j] + gewicht[i] > beste[i]:
                beste[i], vorgaenger[i] = beste[j] + gewicht[i], j
    if not punkte:
        return set()
    kette: set[int] = set()
    i: int | None = max(beste, key=beste.get)
    while i is not None:
        kette.add(i)
        i = vorgaenger[i]
    return {i for i, (_, _, neu) in enumerate(punkte) if neu and i not in kette}


def km_konflikte(punkte: list[tuple[datetime.date, int, str, bool]], max_meldungen: int = 5) -> list[str]:
    """Prueft, dass Kilometerstaende mit dem Datum nicht sinken.

    punkte: (datum, km, bezeichnung, ist_neu). Gemeldet werden nur
    Widersprueche, an denen mindestens ein neuer (importierter) Punkt
    beteiligt ist - Altbestand wurde schon beim Erfassen geprueft.
    """
    meldungen: list[str] = []
    hoechster: tuple[datetime.date, int, str, bool] | None = None
    for punkt in sorted(punkte, key=lambda p: (p[0], p[1])):
        if hoechster is not None and punkt[1] < hoechster[1] and (punkt[3] or hoechster[3]):
            meldungen.append(
                f"{punkt[2]}: {punkt[1]} km am {punkt[0]:%d.%m.%Y} liegt unter "
                f"{hoechster[1]} km vom {hoechster[0]:%d.%m.%Y} ({hoechster[2]})."
            )
            if len(meldungen) >= max_meldungen:
                break
        if hoechster is None or punkt[1] >= hoechster[1]:
            hoechster = punkt
    return meldungen
