"""Excel-Import (Issue #16) mit einer synthetischen Arbeitsmappe, die den
Aufbau der echten Buchfuehrung nachbildet (alle Zahlen frei erfunden)."""
import datetime
import io

import pytest
from openpyxl import Workbook

from app.access import CurrentUser, get_current_user
from app.excel_import import parse_workbook
from app.models import FuelEntry, Kostenkategorie, Kraftstoffart, OtherCost

D = datetime.datetime


def _fuel_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws["A1"] = "Kraftfahrzeug Buchführung"
    ws["E1"] = "Amtl. Kennzeichen"
    ws["G1"] = "TE-ST 1"
    # Monatsuebersicht oben rechts, wie im Original
    for i, monat in enumerate(["Januar", "Februar", "März"]):
        ws.cell(1, 17 + i, monat)
    ws["P2"], ws["Q2"] = 2022, 0
    ws["B3"], ws["C3"] = "Hersteller:", "Testwerke"
    ws["A14"], ws["C14"] = "Gesamtfüllmenge", 999
    headers = ["Datum", "Kilometerstand", "Füllmenge", "Preis/ Liter", "Gesamtpreis", "Ø - Verbrauch",
               "Benzinkosten / Kilometer", "nicht voll", "Benzin", "Kilometer / Liter", "Gesamtfüllmenge"]
    for col, text in enumerate(headers, start=1):
        ws.cell(18, col, text)
    ws.cell(18, 16, "gefahrene KM")
    for r, values in enumerate(rows, start=19):
        for col, value in enumerate(values, start=1):
            if value is not None:
                ws.cell(r, col, value)
        ws.cell(r, 6, "#N/A")  # Hilfsspalte mit Fehlerwert
    ws.cell(19 + len(rows) + 5, 1, "X")  # Endmarkierung
    return ws


def build_workbook(lpg_rows=None, benzin_rows=None, cost_rows=None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _fuel_sheet(wb, "Kraftstoffkosten LPG", lpg_rows if lpg_rows is not None else [
        (D(2022, 3, 1), 100000),                                  # Startzeile (Kauf-km)
        (D(2022, 3, 5), 100300, 40.0, 0.8, 32.0),
        (D(2022, 3, 20), 100500, 10.0, 0.9, 9.0, None, None, 1),  # nicht voll
        (D(2022, 4, 2), 100900, 45.5, 0.8, "=C22*D22"),           # Formel ohne Cache-Wert
        (None, None, None, None, None),                           # Leerzeile
        (D(2022, 4, 10), 101000, 0, 0.8),                         # Nullzeile
        (D(2022, 5, 1), 101400, 42.0, 0.85, 35.7),
    ])
    _fuel_sheet(wb, "Kraftstoffkosten Benzin", benzin_rows if benzin_rows is not None else [
        (D(2022, 3, 1), 100000),
        (D(2022, 4, 2), 100900, 20.0, 1.5, 30.0),  # gleicher Stopp wie LPG-Tankung
    ])

    stat = wb.create_sheet("Gesamtstatistik")
    for col, text in enumerate(["Datum", "Kilometerstand", "Preis/ Liter", "Gesamtpreis", "Hilfsspalte"], start=1):
        stat.cell(3, col, text)
    stat.append([D(2022, 3, 5), 100300, 0.8, 32.0, 1])

    costs = wb.create_sheet("Sonstige Kosten")
    costs["O2"] = 1234.5
    costs["K3"], costs["Q3"] = 99, 4321
    headers = {2: "Datum Kategorie ", 3: "Einnahmen RKA", 4: "Versicherung", 5: "Steuer", 6: "Finanzierung",
               7: "Reifen / Teile", 8: "Service / Tüv", 9: "Waschen", 10: "Parken", 11: "Jährlich",
               12: "Versicherungs-kosten / Tag bis heute", 13: "Steuer bis heute", 14: "Unterhalts- kosten",
               17: "Gesamtkosten"}
    for col, text in headers.items():
        costs.cell(4, col, text)
    rows = cost_rows if cost_rows is not None else [
        # (Datum, {Spalte: Betrag}, Beschreibung, jaehrlich)
        (D(2022, 3, 1), {6: 5000}, "Erwerb", None),
        (D(2022, 3, 1), {5: 30}, "Zulassung", None),
        (D(2022, 3, 10), {4: 300}, "KFZ Versicherung", 1),
        (D(2022, 3, 12), {3: -5.5}, None, None),
        (D(2022, 4, 1), {7: 120}, "Winterreifen", None),
        (D(2022, 4, 3), {9: 8}, None, None),
        (D(2022, 4, 4), {10: 3.5}, "Parkhaus", None),
        (None, {8: 50}, "ohne Datum", None),
    ]
    for r, (datum, amounts, desc, jaehrlich) in enumerate(rows, start=5):
        costs.cell(r, 2, datum)
        for col, betrag in amounts.items():
            costs.cell(r, col, betrag)
            costs.cell(r, 14, betrag)  # berechnete Spalte, darf nicht zaehlen
        if desc:
            costs.cell(r, 16, desc)
        if jaehrlich:
            costs.cell(r, 11, jaehrlich)
    # Formelzeilen ohne Werte am Ende wie im Original
    costs.cell(len(rows) + 6, 17, "=SUM(C20:I20)")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- Parser ------------------------------------------------------------------


def test_parser_liest_tankungen_und_kosten():
    preview = parse_workbook(build_workbook())

    fuel = preview.fuel_entries
    assert [(e.datum.isoformat(), e.kilometerstand, e.kraftstoffart) for e in fuel] == [
        ("2022-03-05", 100300, Kraftstoffart.LPG),
        ("2022-03-20", 100500, Kraftstoffart.LPG),
        ("2022-04-02", 100900, Kraftstoffart.LPG),
        ("2022-04-02", 100900, Kraftstoffart.BENZIN),
        ("2022-05-01", 101400, Kraftstoffart.LPG),
    ]
    assert [e.nicht_voll for e in fuel] == [False, True, False, False, False]
    assert fuel[2].gesamtpreis == pytest.approx(36.4)  # aus Menge x Preis/l ergaenzt
    assert fuel[0].quelle == "Kraftstoffkosten LPG, Zeile 20"

    kosten = {(c.kategorie, c.betrag): c for c in preview.other_costs}
    assert set(kosten) == {
        (Kostenkategorie.STEUER, 30.0),
        (Kostenkategorie.VERSICHERUNG, 300.0),
        (Kostenkategorie.EINNAHME, -5.5),
        (Kostenkategorie.REIFEN_TEILE, 120.0),
        (Kostenkategorie.WASCHEN, 8.0),
        (Kostenkategorie.SONSTIGES, 3.5),
    }
    assert kosten[(Kostenkategorie.VERSICHERUNG, 300.0)].jaehrlich_wiederkehrend is True
    assert kosten[(Kostenkategorie.VERSICHERUNG, 300.0)].beschreibung == "KFZ Versicherung"
    assert kosten[(Kostenkategorie.SONSTIGES, 3.5)].beschreibung == "Parken: Parkhaus"
    assert kosten[(Kostenkategorie.EINNAHME, -5.5)].beschreibung == "Einnahmen RKA"

    # Erwerb wird Kaufvorschlag statt Kosten, Kauf-km aus der Startzeile
    assert preview.kauf.kaufpreis == 5000
    assert preview.kauf.kaufdatum == datetime.date(2022, 3, 1)
    assert preview.kauf.kaufkilometerstand == 100000

    warnungen = " ".join(preview.warnungen)
    assert "Zeile 24" in warnungen  # Nullzeile
    assert "Sonstige Kosten, Zeile 12: kein Datum" in warnungen


def test_parser_lehnt_kaputte_datei_ab():
    from app.excel_import import ExcelImportError

    with pytest.raises(ExcelImportError):
        parse_workbook(b"kein excel")


# --- API ---------------------------------------------------------------------


def _vehicle(client, **extra):
    payload = {"kennzeichen": "TE-ST 1", "hersteller": "Testwerke", "modell": "Van", **extra}
    response = client.post("/api/vehicles", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def _upload(client, vehicle_id, content, dry_run=True):
    return client.post(
        f"/api/vehicles/{vehicle_id}/import/excel?dry_run={'true' if dry_run else 'false'}",
        files={"file": ("buchfuehrung.xlsx", content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_dry_run_duplikat_und_import(client, db_session):
    vehicle = _vehicle(client)
    # bereits von Hand erfasst (Gesamtpreis auf Cent gerundet, km leicht anders)
    manual = client.post("/api/fuel-entries", json={
        "vehicle_id": vehicle["id"], "datum": "2022-03-05", "kilometerstand": 100301, "kraftstoffart": "LPG",
        "fuellmenge_liter": 40, "preis_pro_liter": 0.8, "gesamtpreis": 32.004,
    })
    assert manual.status_code == 201, manual.text
    content = build_workbook()

    response = _upload(client, vehicle["id"], content)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["dry_run"] is True
    assert (body["tankungen_neu"], body["tankungen_duplikate"]) == (4, 1)
    assert (body["kosten_neu"], body["kosten_duplikate"]) == (6, 0)
    assert [t["duplikat"] for t in body["tankungen"]][0] is True
    assert body["kauf"]["uebernommen"] == ["kaufdatum", "kaufkilometerstand", "kaufpreis"]
    assert db_session.query(FuelEntry).count() == 1  # Vorschau schreibt nichts
    assert db_session.query(OtherCost).count() == 0

    response = _upload(client, vehicle["id"], content, dry_run=False)
    assert response.status_code == 200, response.text
    assert response.json()["tankungen_neu"] == 4
    db_session.expire_all()
    assert db_session.query(FuelEntry).count() == 5
    assert db_session.query(OtherCost).count() == 6
    saved = client.get(f"/api/vehicles/{vehicle['id']}").json()
    assert (saved["kaufpreis"], saved["kaufdatum"], saved["kaufkilometerstand"]) == (5000, "2022-03-01", 100000)

    stats = client.get(f"/api/vehicles/{vehicle['id']}/stats").json()
    assert stats["verbrauch_nach_kraftstoff"]  # Voll-zu-Voll klappt mit den Importdaten

    # zweiter Import derselben Datei: alles Duplikate
    body = _upload(client, vehicle["id"], content).json()
    assert (body["tankungen_neu"], body["kosten_neu"]) == (0, 0)
    assert body["kauf"]["uebernommen"] == []


def test_vorhandene_kaufdaten_bleiben(client, db_session):
    vehicle = _vehicle(client, kaufpreis=7777, kaufdatum="2022-02-01")
    response = _upload(client, vehicle["id"], build_workbook(), dry_run=False)
    assert response.status_code == 200, response.text
    assert response.json()["kauf"]["uebernommen"] == ["kaufkilometerstand"]
    saved = client.get(f"/api/vehicles/{vehicle['id']}").json()
    assert (saved["kaufpreis"], saved["kaufdatum"]) == (7777, "2022-02-01")


def test_einzelner_tippfehler_im_km_stand_wird_uebersprungen(client, db_session):
    """Wie in der echten Datei: ein km-Stand mit Zahlendreher (181.901 statt
    172.901) darf nicht den ganzen Import blockieren."""
    vehicle = _vehicle(client)
    content = build_workbook(lpg_rows=[
        (D(2022, 3, 5), 100300, 40.0, 0.8, 32.0),
        (D(2022, 3, 20), 190700, 40.0, 0.8, 32.1),  # Tippfehler
        (D(2022, 4, 5), 100700, 40.0, 0.8, 32.2),
        (D(2022, 4, 25), 101100, 40.0, 0.8, 32.3),
    ], benzin_rows=[])
    body = _upload(client, vehicle["id"], content, dry_run=False).json()
    assert body["tankungen_uebersprungen"] == 1
    assert [t["kilometerstand"] for t in body["tankungen"] if t["uebersprungen"]] == [190700]
    assert any("190.700 km" in w and "Tippfehler" in w for w in body["warnungen"])
    gespeichert = sorted(e.kilometerstand for e in db_session.query(FuelEntry).all())
    assert 190700 not in gespeichert and 101100 in gespeichert


def test_viele_sinkende_kilometerstaende_werden_abgelehnt(client, db_session):
    vehicle = _vehicle(client)
    content = build_workbook(lpg_rows=[
        (D(2022, 3, 5), 100300, 40.0, 0.8, 32.0),
        (D(2022, 3, 10), 100200, 40.0, 0.8, 32.1),
        (D(2022, 3, 15), 100100, 40.0, 0.8, 32.2),
        (D(2022, 3, 20), 100000, 40.0, 0.8, 32.3),
    ])
    for dry_run in (True, False):
        response = _upload(client, vehicle["id"], content, dry_run=dry_run)
        assert response.status_code == 400
        assert "Kilometerstände sinken" in response.json()["detail"]
    assert db_session.query(FuelEntry).count() == 0
    assert db_session.query(OtherCost).count() == 0


def test_widerspruch_zu_vorhandener_tankung_wird_abgelehnt(client, db_session):
    vehicle = _vehicle(client)
    client.post("/api/fuel-entries", json={
        "vehicle_id": vehicle["id"], "datum": "2022-03-10", "kilometerstand": 200000, "kraftstoffart": "LPG",
        "fuellmenge_liter": 40, "preis_pro_liter": 0.8, "gesamtpreis": 50,
    })
    response = _upload(client, vehicle["id"], build_workbook())
    assert response.status_code == 400
    assert "200000" in response.json()["detail"]


def test_ungueltige_datei(client, db_session):
    vehicle = _vehicle(client)
    response = _upload(client, vehicle["id"], b"das ist keine Arbeitsmappe")
    assert response.status_code == 400
    assert ".xlsx" in response.json()["detail"]


def test_fremder_nutzer_bekommt_404(client, db_session):
    from main import app

    app.dependency_overrides[get_current_user] = lambda: CurrentUser("alice")
    try:
        vehicle = _vehicle(client)
        app.dependency_overrides[get_current_user] = lambda: CurrentUser("carol")
        response = _upload(client, vehicle["id"], build_workbook(), dry_run=False)
        assert response.status_code == 404
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert db_session.query(FuelEntry).count() == 0
